"""Gera o SQL que carrega a aba BD da planilha BD_Talentos.xlsx nas tabelas do Supabase.

Uso: python scripts/importar_planilha.py <planilha.xlsx> <saida.sql>

Pode rodar de novo com a planilha atualizada: o SQL apaga as vagas que vieram
da planilha e recarrega, sem mexer nas vagas criadas pelo sistema.
"""

import datetime as dt
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict

import openpyxl

COLUNAS = [
    "Projeto", "Empresa Faturamento", "Abertura", "Cliente", "Cidade", "Vaga",
    "Receita", "Analista", "Status", "OBS", "Data de conclusão", "Faturamento",
]

# Mesmo cliente escrito de formas que a normalização automática não junta.
CLIENTES_MESMO_NOME = {
    "Punch Sabia": "Punch - Patio Sabia",
    "Punch Sabiá": "Punch - Patio Sabia",
    "Punch Vinhedos": "Punch - Patio Vinhedos",
    "Akkar Sabiá": "Akkar Patio Sabia",
    "Casa da Parmegiana": "Casa do Parmegiana",
    "Gellato Borelli": "Gelato Borelli",
    "Gelato Uberlandia Shopping": "Gelato Borelli - Uberlândia Shopping",
    "Gelato Landscape": "Gelato Borelli - Landscape",
    "Santos EPI's": "Santos EPIs",
    "BusinessCont": "Business Cont",
    "Ecos Solar": "Ecos Energia Solar",
    "Morada Nova - Mat Constr": "Morada Nova (Mat. Construção)",
    "Digital + cont": "Digital + Contabilidade",
    "Contabilidade Consultec": "Consultec Contabilidade",
    "Contabilidade Hoffmam": "Hoffmam Contabilidade",
}

# Datas digitadas com erro na planilha (conferidas uma a uma).
DATAS_CORRIGIDAS = {
    "26/09/20/25": dt.date(2025, 9, 26),
    "05/0622025": dt.date(2025, 6, 5),
    "17/062025": dt.date(2025, 6, 17),
    "29/07/0205": dt.date(2025, 7, 29),
    "21/092/2026": dt.date(2026, 9, 21),
}

STATUS = {
    "concluída": "concluida",
    "cancelada": "cancelada",
    "aberta": "aberta",
    "substituição": "substituicao",
    "faturar": "faturar",
    "congelada": "congelada",
}

SUFIXO_CONTABIL = re.compile(r"\s+(contabilidade|contabil|contábil)$", re.IGNORECASE)


def chave(texto):
    sem_acento = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", sem_acento.lower()).strip()


def limpo(valor):
    if valor is None:
        return None
    texto = re.sub(r"\s+", " ", str(valor)).strip()
    return texto or None


def escolher_grafia(variantes, preferidos):
    """Nome da lista oficial, depois com inicial maiúscula, depois o mais usado, depois com acento."""
    return max(variantes, key=lambda v: (
        v in preferidos, v[:1].isupper(), variantes[v], sum(ord(c) > 127 for c in v), len(v),
    ))


def unificar(nomes, chave_de, preferidos=frozenset()):
    grupos = defaultdict(Counter)
    for nome in nomes:
        grupos[chave_de(nome)][nome] += 1
    canonico = {}
    for variantes in grupos.values():
        escolhido = escolher_grafia(variantes, preferidos)
        for v in variantes:
            canonico[v] = escolhido
    return canonico


def chave_cliente(nome):
    base = chave(SUFIXO_CONTABIL.sub("", nome))
    return base or chave(nome)


def data(valor, avisos, linha, coluna):
    if valor is None:
        return None, None
    if isinstance(valor, dt.datetime):
        return valor.date(), None
    texto = limpo(valor)
    if texto in DATAS_CORRIGIDAS:
        avisos.append(f"linha {linha}: {coluna} '{texto}' corrigida para {DATAS_CORRIGIDAS[texto]:%d/%m/%Y}")
        return DATAS_CORRIGIDAS[texto], None
    return None, texto


def tipo_faturamento(texto):
    if not texto or texto == "-":
        return None
    t = texto.lower()
    if t.startswith("ok"):
        return "substituicao" if "subst" in t else "normal"
    if t.startswith(("sub", "subt")):
        return "substituicao"
    if t.startswith("repos"):
        return "reposicao"
    if t.startswith("bon"):
        return "bonificacao"
    if t.startswith("cancel"):
        return "cancelada"
    raise ValueError(f"Faturamento desconhecido: {texto!r}")


def projeto_e_canal(projeto):
    projeto = limpo(projeto)
    if projeto.startswith("Escalar"):
        _, _, canal = projeto.partition(" - ")
        return "Escalar", canal or None
    if chave(projeto) == "outros negocios":
        return "Outros Negócios", None
    return projeto, None


def empresa_e_canal(empresa, projeto, avisos, linha):
    empresa = limpo(empresa)
    if empresa and empresa.startswith("Escalar Talentos"):
        _, _, canal = empresa.partition(" - ")
        return "Escalar Talentos", canal or None, False
    if empresa == "Escritorial Talentos":
        return empresa, None, False
    if empresa:
        avisos.append(f"linha {linha}: empresa de faturamento '{empresa}' não existe; preenchida pelo projeto")
    return ("Escalar Talentos" if projeto == "Escalar" else "Escritorial Talentos"), None, True


def cargo_e_quantidade(vaga):
    vaga = limpo(vaga)
    achou = re.search(r"\((\d{1,2})\)", vaga)
    if not achou:
        return vaga, 1
    cargo = limpo(vaga[: achou.start()] + vaga[achou.end():])
    return cargo, int(achou.group(1))


def sql(valor):
    if valor is None:
        return "null"
    if isinstance(valor, (int, float)):
        return repr(valor)
    if isinstance(valor, dt.date):
        return f"'{valor.isoformat()}'"
    return "'" + str(valor).replace("'", "''") + "'"


def clientes_oficiais(caminho):
    livro = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    if "Página5" not in livro.sheetnames:
        return frozenset()
    valores = (limpo(c) for linha in livro["Página5"].iter_rows(values_only=True) for c in linha)
    return frozenset(v for v in valores if v and v != "Cliente")


def ler_linhas(caminho):
    aba = openpyxl.load_workbook(caminho, data_only=True)["BD"]
    cabecalho = [limpo(c) for c in next(aba.iter_rows(min_row=1, max_row=1, values_only=True))]
    faltando = [c for c in COLUNAS if c not in cabecalho]
    if faltando:
        sys.exit(f"Colunas não encontradas na aba BD: {faltando}")
    idx = {c: cabecalho.index(c) for c in COLUNAS}
    for numero, valores in enumerate(aba.iter_rows(min_row=2, values_only=True), start=2):
        linha = {c: valores[i] for c, i in idx.items()}
        if any(v is not None and str(v).strip() for v in linha.values()):
            yield numero, linha


def main(entrada, saida):
    linhas = list(ler_linhas(entrada))
    avisos = []

    nomes_clientes = [CLIENTES_MESMO_NOME.get(limpo(l["Cliente"]), limpo(l["Cliente"])) for _, l in linhas]
    cliente_canonico = unificar(nomes_clientes, chave_cliente, clientes_oficiais(entrada))
    cidade_canonica = unificar([limpo(l["Cidade"]) for _, l in linhas if limpo(l["Cidade"])], chave)

    vagas, empresas_inferidas = [], 0
    for numero, l in linhas:
        projeto, canal = projeto_e_canal(l["Projeto"])
        empresa, canal_empresa, inferida = empresa_e_canal(l["Empresa Faturamento"], projeto, avisos, numero)
        empresas_inferidas += inferida
        cargo, quantidade = cargo_e_quantidade(l["Vaga"])
        abertura, abertura_texto = data(l["Abertura"], avisos, numero, "Abertura")
        if abertura is None:
            sys.exit(f"linha {numero}: data de abertura inválida {abertura_texto!r}")
        conclusao, conclusao_texto = data(l["Data de conclusão"], avisos, numero, "Conclusão")
        observacoes = limpo(l["OBS"])
        if conclusao_texto and conclusao_texto != "-":
            observacoes = f"{observacoes} (conclusão: {conclusao_texto})" if observacoes else f"Conclusão: {conclusao_texto}"
        status = STATUS.get(limpo(l["Status"]).lower())
        if status is None:
            sys.exit(f"linha {numero}: status desconhecido {l['Status']!r}")
        nome_cliente = CLIENTES_MESMO_NOME.get(limpo(l["Cliente"]), limpo(l["Cliente"]))
        original = {c: (v.isoformat() if isinstance(v, dt.datetime) else v) for c, v in l.items()}
        vagas.append([
            numero, projeto, empresa, canal or canal_empresa, cliente_canonico[nome_cliente],
            cidade_canonica.get(limpo(l["Cidade"])), cargo, quantidade,
            round(float(l["Receita"] or 0), 2), limpo(l["Analista"]), status,
            tipo_faturamento(limpo(l["Faturamento"])), observacoes, abertura, conclusao,
            json.dumps(original, ensure_ascii=False),
        ])

    analistas = sorted({v[9] for v in vagas if v[9]})
    clientes = sorted({v[4] for v in vagas})

    with open(saida, "w", encoding="utf-8") as f:
        f.write("-- Gerado por scripts/importar_planilha.py a partir da aba BD.\nbegin;\n\n")
        f.write("insert into public.analistas (nome) values\n")
        f.write(",\n".join(f"  ({sql(a)})" for a in analistas) + "\non conflict (nome) do nothing;\n\n")
        f.write("insert into public.clientes (nome) values\n")
        f.write(",\n".join(f"  ({sql(c)})" for c in clientes) + "\non conflict (nome) do nothing;\n\n")
        f.write("delete from public.vagas where planilha_linha is not null;\n\n")
        f.write(
            "insert into public.vagas (planilha_linha, projeto_id, empresa_faturamento, canal, cliente_id, cidade,\n"
            "  cargo, quantidade, receita, analista_id, status, tipo_faturamento, observacoes,\n"
            "  data_abertura, data_conclusao, planilha_original)\n"
            "select t.linha::int, p.id, t.empresa, t.canal, c.id, t.cidade, t.cargo, t.quantidade::int,\n"
            "  t.receita::numeric, a.id, t.status, t.tipo, t.obs, t.abertura::date, t.conclusao::date, t.original::jsonb\n"
            "from (values\n"
        )
        f.write(",\n".join("  (" + ", ".join(sql(x) for x in v) + ")" for v in vagas))
        f.write(
            "\n) as t(linha, projeto, empresa, canal, cliente, cidade, cargo, quantidade, receita,\n"
            "        analista, status, tipo, obs, abertura, conclusao, original)\n"
            "join public.projetos p on p.nome = t.projeto\n"
            "join public.clientes c on c.nome = t.cliente\n"
            "left join public.analistas a on a.nome = t.analista;\n\n"
        )
        f.write(f"do $$ begin\n  if (select count(*) from public.vagas where planilha_linha is not null) <> {len(vagas)} then\n"
                f"    raise exception 'Importação incompleta: esperava {len(vagas)} vagas';\n  end if;\nend $$;\n\ncommit;\n")

    juntados = defaultdict(set)
    for original, destino in cliente_canonico.items():
        if original != destino:
            juntados[destino].add(original)
    presentes = {limpo(l["Cliente"]) for _, l in linhas}
    for original, destino in CLIENTES_MESMO_NOME.items():
        if original in presentes:
            juntados[cliente_canonico.get(destino, destino)].add(original)

    print(f"Vagas: {len(vagas)} | Clientes: {len(clientes)} | Analistas: {len(analistas)} ({', '.join(analistas)})")
    print(f"Receita total: R$ {sum(v[8] for v in vagas):,.2f}")
    print(f"Empresa de faturamento preenchida pelo projeto: {empresas_inferidas} vagas")
    print(f"Clientes unificados ({len(juntados)}):")
    for destino, origens in sorted(juntados.items()):
        print(f"  {destino}  <-  {', '.join(sorted(origens))}")
    print("Avisos:")
    for a in avisos:
        print(" ", a)
    print(f"SQL gravado em {saida}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
